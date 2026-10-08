"""Turn a cluster table into a color-coded Excel workbook for people to review the clusters.

Sheets (colors match the 3D pages):
  Overview  one line per cluster: color, theme words, number of notes
  Clusters  every clustered note, grouped by cluster, with its source (German original, post-it, scan)
  Noise     the notes that fit no cluster
Each note sheet has an empty `comment` column for the reviewer.

Called by src/embed_cluster.py; also runs alone on the saved CSVs:
    python src/cluster_sheet.py
"""

import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from plot_html import NOISE_COLOR, OUT_DIR, PALETTE, SCAN_DIR, source

HEADER_FILL = PatternFill("solid", fgColor="1D2240")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="D0D4E0")
WRAP = Alignment(wrap_text=True, vertical="top")
COLUMNS = [("cluster", 9), ("theme", 30), ("note", 52), ("original (German)", 42), ("post-it", 10),
           ("side/row/col/part", 16), ("scan", 24), ("note #", 8), ("checked", 18), ("comment", 30)]


def to_hex(color):
    """'rgb(r, g, b)' or '#rrggbb' -> 'RRGGBB'."""
    if color.startswith("#"):
        return color[1:].upper()
    return "".join(f"{int(v):02X}" for v in re.findall(r"\d+", color)[:3])


def tint(hex_color, keep=0.25):
    """A light version of a color for row backgrounds: `keep` of the color, the rest white."""
    rgb = [int(hex_color[i:i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(255 - (255 - v) * keep):02X}" for v in rgb)


def cluster_colors(df):
    """The same color per cluster as the 3D pages: palette order over the sorted cluster ids."""
    ids = sorted(c for c in df["cluster"].unique() if c != -1)
    colors = {c: to_hex(PALETTE[i % len(PALETTE)]) for i, c in enumerate(ids)}
    colors[-1] = to_hex(NOISE_COLOR)
    return colors


def header(ws, columns):
    for k, (name, width) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=k, value=name)
        cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(k)].width = width
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"


def note_sheet(ws, df, colors):
    header(ws, COLUMNS)
    for r, (_, row) in enumerate(df.iterrows(), 2):
        n = source(int(row["row"]))
        original = n["original"] if n["original"].strip() != n["english"].strip() else ""
        values = ["noise" if row["cluster"] == -1 else int(row["cluster"]), row["theme"], n["text"], original,
                  n["where"], f"{n['seite']} / {n['zeile']} / {n['spalte']} / {n['teil']}", n["scan"], n["row"],
                  n["flags"], ""]
        fill = PatternFill("solid", fgColor=tint(colors[row["cluster"]]))
        for k, value in enumerate(values, 1):
            cell = ws.cell(row=r, column=k, value=value)
            cell.fill, cell.alignment, cell.border = fill, WRAP, Border(bottom=THIN)
        # a solid color chip in the cluster column, like the dot in the plot
        chip = ws.cell(row=r, column=1)
        chip.fill = PatternFill("solid", fgColor=colors[row["cluster"]])
        chip.font, chip.alignment = Font(bold=True, color="FFFFFF"), Alignment(horizontal="center", vertical="top")
        scan = SCAN_DIR / "Worshop scan" / n["scan"]
        if scan.exists():
            ws.cell(row=r, column=7).hyperlink = str(scan)
            ws.cell(row=r, column=7).font = Font(color="1F4FD8", underline="single")
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(len(df) + 1, 2)}"


def write_sheet(df, path):
    colors = cluster_colors(df)
    wb = Workbook()
    ws = wb.active
    ws.title = "Overview"
    header(ws, [("cluster", 9), ("theme words", 46), ("notes", 8), ("share", 8)])
    counts = df["cluster"].value_counts()
    for r, c in enumerate(sorted(counts.index, key=lambda c: (c == -1, c)), 2):
        theme = "noise: fits no cluster" if c == -1 else df.loc[df["cluster"] == c, "theme"].iloc[0]
        for k, value in enumerate(["noise" if c == -1 else int(c), theme, int(counts[c]), counts[c] / len(df)], 1):
            cell = ws.cell(row=r, column=k, value=value)
            cell.fill, cell.border = PatternFill("solid", fgColor=tint(colors[c])), Border(bottom=THIN)
        ws.cell(row=r, column=4).number_format = "0%"
        chip = ws.cell(row=r, column=1)
        chip.fill, chip.font = PatternFill("solid", fgColor=colors[c]), Font(bold=True, color="FFFFFF")
        chip.alignment = Alignment(horizontal="center")
    total = len(counts) + 2
    ws.cell(row=total, column=2, value="total").font = Font(bold=True)
    ws.cell(row=total, column=3, value=len(df)).font = Font(bold=True)

    clustered = df[df["cluster"] != -1].sort_values(["cluster", "row"])
    note_sheet(wb.create_sheet("Clusters"), clustered, colors)
    note_sheet(wb.create_sheet("Noise"), df[df["cluster"] == -1].sort_values("row"), colors)
    wb.save(path)


def main():
    paths = [Path(p) for p in sys.argv[1:]] or sorted(OUT_DIR.glob("clusters_*.csv"))
    for csv in paths:
        write_sheet(pd.read_csv(csv, encoding="utf-8-sig"), csv.with_suffix(".xlsx"))
        print(f"{csv.name} -> {csv.stem}.xlsx")


if __name__ == "__main__":
    main()
